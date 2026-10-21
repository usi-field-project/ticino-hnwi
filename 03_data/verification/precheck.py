"""Automatic pre-check of observations.csv against the local source files.

For each row, looks for the recorded value among the numbers in `local_file`
(xlsx, csv, pdf, docx). It does NOT verify anything: a value becomes
`verified` only after a second team member has opened the source
(03_data/README.md). The output tells the reviewer where to look first.

    python 03_data/verification/precheck.py
"""
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
import re
import subprocess

import warnings

import pandas as pd

warnings.filterwarnings("ignore", module="openpyxl")

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).with_name("precheck.csv")
NUM = re.compile(r"(?<![\w.])-?\d{1,3}(?:\.\d{3}){2,}(?:,\d+)?(?![\d.])|(?<![\w.])-?\d{1,3}(?:[’'  ,]\d{3})+(?:[.,]\d+)?(?![\d])|(?<![\w.])-?\d+(?:[.,]\d+)?(?![\d])")


def parse(token):
    if token.count(".") > 1:  # Italian/German thousands separator: 1.407.044
        token = token.replace(".", "").replace(",", ".")
    t = token.replace("’", "").replace("'", "").replace(" ", "").replace(" ", "")
    out = set()
    if "," in t and "." in t:
        out.add(t.replace(",", ""))
    elif "," in t:
        head, tail = t.rsplit(",", 1)
        out.add(t.replace(",", ""))          # thousands separator
        out.add(head.replace(",", "") + "." + tail)  # decimal comma
    else:
        out.add(t)
    vals = set()
    for s in out:
        try:
            vals.add(float(s))
        except ValueError:
            pass
    return vals


@lru_cache(maxsize=None)
def numbers_in(path):
    p = ROOT / path
    ext = p.suffix.lower()
    nums, texts = set(), []
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        for ws in wb.worksheets:
            ws.reset_dimensions()  # some exports declare a wrong sheet size
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if isinstance(c, (int, float)) and not isinstance(c, bool):
                        nums.add(float(c))
                    elif isinstance(c, str):
                        texts.append(c)
    elif ext == ".csv":
        import csv
        raw = p.read_text(encoding="utf-8-sig", errors="ignore")
        delim = ";" if raw.count(";") > raw.count(",") else ","
        texts += [cell for row in csv.reader(raw.splitlines(), delimiter=delim) for cell in row]
    elif ext == ".pdf":
        texts.append(subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout)
    elif ext == ".docx":
        import docx
        d = docx.Document(p)
        texts += [x.text for x in d.paragraphs]
        texts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    for t in texts:
        for tok in NUM.findall(t):
            nums |= parse(tok)
    return frozenset(nums)


def decimals(v):
    e = Decimal(v).as_tuple().exponent
    return max(0, -e)


def matches(v, src):
    x = float(v)
    d = decimals(v)
    tol = 0.5 * 10 ** -d + 1e-9
    for scale, label in [(1, "exact"), (100, "x100"), (0.01, "/100"), (1e-3, "/1000"), (1e3, "x1000"), (1e-6, "/1e6"), (1e6, "x1e6")]:
        if any(abs(s * scale - x) <= tol for s in src):
            return label
    return ""


def arithmetic(v, nums):
    """Return (operation, inputs) if v is a sum, difference or ratio of figures in the note."""
    from itertools import combinations
    x, tol = float(v), 0.5 * 10 ** -decimals(v) + 1e-9
    nums = sorted((n for n in nums if n != 0), reverse=True)[:12]
    for k in (2, 3, 4):
        for c in combinations(nums, k):
            if abs(sum(c) - x) <= tol:
                return "sum", list(c)
    for a in nums:
        for base in (1e6, 1e7):  # tax as % of the reference income or wealth
            if abs(a / base * 100 - x) <= tol:
                return "share of the reference base", [a]
        for b in nums:
            if a is b:
                continue
            if abs(a - b - x) <= tol:
                return "difference", [a, b]
            for k in (1, 100, 1000, 1e5):
                if abs(a / b * k - x) <= tol:
                    return "ratio", [a, b]
    return None


def recompute(r):
    """Recalculate values that are sums or rates over a whole source table."""
    f = ROOT / r.local_file
    label = r.source_indicator_label
    canton = {"ZH": "Zürich", "TI": "Ticino", "GE": "Genève", "GR": "Graubünden", "ZG": "Zug"}.get(r.geo_id)
    if canton is None:
        return None
    if f.name.startswith("kzp24_data"):  # FOPH hospital key figures: sum of bed columns by canton
        d = pd.read_excel(f, sheet_name="KZ2024_KZP24")
        cols = [c for c in d.columns if c.startswith("BettenStat")]
        val = d[d.KT == r.geo_id][cols].apply(pd.to_numeric, errors="coerce").sum().sum()
        return round(val, 1), f"sum of {', '.join(cols)} over {len(d[d.KT == r.geo_id])} hospitals"
    if "STATENT" in f.name and "NOGA 64" in label:
        d = pd.read_csv(f)
        d = d[(d.anno == int(r.period)) & d.cantone.str.startswith(canton) & d.settore.str.match(r"^6[456] ")]
        return float(pd.to_numeric(d["Beschäftigte"], errors="coerce").sum()), "sum of NOGA 64, 65, 66 (Beschäftigte)"
    if f.name.startswith("su-i-40.02.03.02.03") and "Erwerbstätigenquote" in label:
        d = pd.read_excel(f, sheet_name=0, header=None)
        row = d[d[2].astype(str).str.startswith(canton.split()[0])].iloc[0]
        pop, active, unemployed, published = row[9], row[11], row[13], row[15]
        calc = (active - unemployed) / pop * 100
        return round(calc, 1), f"(col. 'Popolazione attiva' − 'Persone disoccupate') / 'Popolazione totale', 15-64; the file also publishes 'Tasso di occupazione netto' = {published:.1f}%"
    return None


def check(r):
    if r.value in ("yes", "no"):
        return "manual", "qualitative value (yes/no): read the source"
    if float(r.value) == 0:
        return "manual", "zero (e.g. no tax): read the source"
    rc = recompute(r)
    if rc is not None:
        val, how = rc
        ok = abs(val - float(r.value)) <= 0.5 * 10 ** -decimals(r.value) + 1e-9
        return ("recomputed_ok" if ok else "recomputed_diff"), f"recomputed from the file: {val:g} = {how}"
    src = numbers_in(r.local_file)
    if not src:
        return "no_text", "no numbers extracted from the file (scanned PDF or image?): check manually"
    m = matches(r.value, src)
    sig = len(r.value.replace("-", "").replace(".", "").strip("0"))
    if m == "exact":
        if sig <= 2:
            return "weak", "found in the file, but a short number matches easily: confirm the location"
        return "found", ""
    note_nums = []
    text = re.sub(r"\b1[56]\s*[-–]\s*(64|74)\b", " ", r.comparability_note + " " + r.source_indicator_label)  # age ranges
    for tok in NUM.findall(text):
        note_nums += sorted(parse(tok), reverse=True)
    how = arithmetic(r.value, note_nums)
    if how:
        found = [n for n in how[1] if n in src]
        status = "derived_ok" if len(found) == len(how[1]) else "derived_partial"
        return status, f"value = {how[0]} of figures in the note; {len(found)} of {len(how[1])} figures found in the file"
    if m:
        return "found_scaled", f"found written in another unit ({m}, e.g. billion vs million): check the unit"
    if any(n in src for n in note_nums if abs(n) >= 100):
        return "derived", "value not in the file, but figures quoted in the note are: probably calculated, check the calculation"
    return "not_found", "value not found in the file: check value, unit and file"


obs = pd.read_csv(ROOT / "03_data/observations.csv", dtype=str).fillna("")
res = obs.apply(lambda r: pd.Series(check(r), index=["precheck", "precheck_note"]), axis=1)
out = pd.concat([obs[["obs_id", "indicator_id", "value", "unit", "collected_by", "local_file"]], res], axis=1)
out.to_csv(OUT, index=False)
print(out.precheck.value_counts().to_string())
